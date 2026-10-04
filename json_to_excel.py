#!/usr/bin/env python3
import json, glob, os, re, sys
from pathlib import Path
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

_VER_NUM_RE = re.compile(r"\d+")


def _version_sort_key(v):
    parts = _VER_NUM_RE.findall(str(v or ""))
    if parts:
        return (0, tuple(int(x) for x in parts[:6]), "")
    return (1, (), str(v or "").lower())


def _merge_uniq_by_name(uniq):
    """Group (name, version) entries by software name (case-insensitive).

    One row per Software; Version cell holds all distinct versions sorted.
    Count = total installs across versions; hosts = union preserving order.
    """
    groups = OrderedDict()
    for v in (uniq or {}).values():
        key = str(v.get("name") or "").strip().lower()
        if not key:
            continue
        g = groups.get(key)
        if g is None:
            g = groups[key] = {
                "name": str(v.get("name") or "").strip(),
                "versions": [], "versions_set": set(),
                "publishers": [], "pub_set": set(),
                "count": 0, "hosts": [], "hosts_set": set(),
            }
        ver = str(v.get("version") or "").strip()
        if ver and ver not in g["versions_set"]:
            g["versions_set"].add(ver)
            g["versions"].append(ver)
        pub = str(v.get("publisher") or "").strip()
        if pub and pub not in g["pub_set"]:
            g["pub_set"].add(pub)
            g["publishers"].append(pub)
        try:
            g["count"] += int(v.get("count", 0))
        except (TypeError, ValueError):
            pass
        for h in (v.get("hosts") or []):
            if h not in g["hosts_set"]:
                g["hosts_set"].add(h)
                g["hosts"].append(h)
    merged = OrderedDict()
    for key, g in groups.items():
        g["versions"].sort(key=_version_sort_key)
        merged[key] = {
            "name": g["name"],
            "version": ", ".join(g["versions"]),
            "publisher": " / ".join(g["publishers"]),
            "count": g["count"],
            "hosts": g["hosts"],
        }
    return OrderedDict(sorted(merged.items(), key=lambda kv: kv[0]))

def get_app_dir():
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).parent
    return Path(__file__).parent

def get_out_dir():
    d = get_app_dir() / "out"
    d.mkdir(parents=True, exist_ok=True)
    return d

def find_all(pattern="kaspersky_inventory_*.json"):
    base = get_app_dir()
    files = glob.glob(str(base / pattern)) + glob.glob(str(base / "out" / pattern))
    files = sorted(set(files), key=os.path.getmtime)
    return files

def load_json(p): return json.loads(Path(p).read_text(encoding="utf-8"))

def classify_os(os_name):
    """Windows / Linux / Other from the device os field text."""
    s = str(os_name or "").strip().lower()
    if not s:
        return "other"
    if "windows" in s or "microsoft" in s or "win32" in s:
        return "windows"
    linux_keywords = (
        "linux", "ubuntu", "debian", "centos", "rhel", "red hat", "fedora",
        "suse", "astra", "alt linux", "rosa", "mint", "arch", "gentoo",
        "oracle linux", "almalinux", "rocky", "kali",
    )
    if any(k in s for k in linux_keywords):
        return "linux"
    return "other"


def dedup_by_os(data):
    """Unique software per OS category + host count per category."""
    split = {"windows": OrderedDict(), "linux": OrderedDict(), "other": OrderedDict()}
    host_counts = {"windows": 0, "linux": 0, "other": 0}
    for h in data.get("hosts", []):
        host_info = h.get("host", {}) if isinstance(h.get("host"), dict) else {}
        os_name = host_info.get("os", "")
        ip_raw = str(host_info.get("ip") or "").strip()
        # Hosts sheets show IP instead of hostname (fallback to name when IP missing)
        fallback = host_info.get("display_name") or host_info.get("dn") or "unknown"
        hid = ip_raw.split(",")[0].strip() if ip_raw else str(fallback).strip()
        if not hid:
            hid = "unknown"
        cat = classify_os(os_name)
        host_counts[cat] += 1
        uniq = split[cat]
        for sw in h.get("software", []):
            if not isinstance(sw, dict):
                continue
            name = str(sw.get("name") or "").strip()
            ver = str(sw.get("version") or "").strip()
            pub = str(sw.get("publisher") or "").strip()
            if not name:
                continue
            key = (name.lower(), ver)
            if key not in uniq:
                uniq[key] = {"name": name, "version": ver, "publisher": pub, "count": 1, "hosts": [hid]}
            else:
                uniq[key]["count"] += 1
                if hid not in uniq[key]["hosts"]:
                    uniq[key]["hosts"].append(hid)
    for cat in split:
        split[cat] = OrderedDict(sorted(split[cat].items(), key=lambda kv: (kv[0][0], kv[0][1])))
    return split, host_counts

def _write_software_sheet(ws, uniq, hdr_font, hdr_fill, bd, center, left):
    headers = ["Software Name", "Version", "Publisher", "Hosts Count"]
    ws.append(headers)
    for c in range(1, 5):
        x = ws.cell(row=1, column=c); x.font = hdr_font; x.fill = hdr_fill; x.alignment = center; x.border = bd
    r = 2
    for v in uniq.values():
        ws.cell(row=r, column=1, value=v["name"]).alignment = left
        ws.cell(row=r, column=2, value=v["version"]).alignment = center
        ws.cell(row=r, column=3, value=v["publisher"]).alignment = left
        ws.cell(row=r, column=4, value=v["count"]).alignment = center
        for col in range(1, 5):
            ws.cell(row=r, column=col).border = bd
        r += 1
    ws.column_dimensions["A"].width = 55; ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 28; ws.column_dimensions["D"].width = 12
    ws.freeze_panes = "A2"
    if r > 2:
        ws.auto_filter.ref = f"A1:D{r-1}"
    return r

def _write_hosts_sheet(ws, uniq, hdr_font, hdr_fill, bd, center, left):
    """Row-based mapping (Software, Version, Host) to avoid Excel 16384-column limit.

    Old layout used one column per software and broke with thousands of
    unique products. New layout scales to ~1M rows.
    """
    headers = ["Software Name", "Version", "Host (IP/name)"]
    ws.append(headers)
    for c in range(1, 4):
        x = ws.cell(row=1, column=c)
        x.font = hdr_font
        x.fill = hdr_fill
        x.alignment = center
        x.border = bd
    r = 2
    for v in uniq.values():
        name = v["name"]
        ver = v.get("version", "")
        hosts = v.get("hosts", [])
        if not hosts:
            ws.cell(row=r, column=1, value=name).alignment = left
            ws.cell(row=r, column=2, value=ver).alignment = center
            ws.cell(row=r, column=3, value="").alignment = left
            for col in range(1, 4):
                ws.cell(row=r, column=col).border = bd
            r += 1
            continue
        for h in hosts:
            ws.cell(row=r, column=1, value=name).alignment = left
            ws.cell(row=r, column=2, value=ver).alignment = center
            ws.cell(row=r, column=3, value=h).alignment = left
            for col in range(1, 4):
                ws.cell(row=r, column=col).border = bd
            r += 1
            if r > 1048575:  # Excel row limit guard
                break
        if r > 1048575:
            break
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 30
    ws.freeze_panes = "A2"
    if r > 2:
        ws.auto_filter.ref = f"A1:C{r-1}"

def write_excel(split_or_uniq, out, host_counts=None):
    """Write Excel file split by OS: Windows / Linux / Other.

    Accepts a split dict {"windows":..., "linux":..., "other":...}
    or (legacy) a merged uniq dict placed in a single sheet.
    """
    # legacy: merged uniq, not a split
    if split_or_uniq and all(isinstance(v, dict) and "name" in v for v in split_or_uniq.values()):
        split = {"windows": split_or_uniq, "linux": OrderedDict(), "other": OrderedDict()}
        host_counts = host_counts or {}
    else:
        split = split_or_uniq
    host_counts = host_counts or {}
    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin = Side(style="thin", color="B4C6E7"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Software sheets show one row per Software (versions merged in Version cell).
    merged = {cat: _merge_uniq_by_name(split.get(cat, OrderedDict()))
              for cat in ("windows", "linux", "other")}

    # Sheet 1: Summary
    ws0 = wb.active; ws0.title = "Summary"
    ws0.append(["OS Type", "Hosts Count", "Unique Software Count"])
    for c in range(1, 4):
        x = ws0.cell(row=1, column=c); x.font = hdr_font; x.fill = hdr_fill; x.alignment = center; x.border = bd
    r = 2
    for cat, label in (("windows", "Windows"), ("linux", "Linux"), ("other", "Other")):
        ws0.cell(row=r, column=1, value=label).alignment = center
        ws0.cell(row=r, column=2, value=host_counts.get(cat, 0)).alignment = center
        ws0.cell(row=r, column=3, value=len(merged.get(cat, OrderedDict()))).alignment = center
        for col in range(1, 4):
            ws0.cell(row=r, column=col).border = bd
        r += 1
    ws0.column_dimensions["A"].width = 18; ws0.column_dimensions["B"].width = 14
    ws0.column_dimensions["C"].width = 22
    ws0.freeze_panes = "A2"

    # Sheets 2-4: Software per OS (merged: one row per Software)
    for cat, title in (("windows", "Windows Software"), ("linux", "Linux Software"), ("other", "Other Software")):
        ws = wb.create_sheet(title)
        _write_software_sheet(ws, merged.get(cat, OrderedDict()), hdr_font, hdr_fill, bd, center, left)

    # Sheets 5-7: Hosts per Software per OS
    for cat, title in (("windows", "Hosts - Windows"), ("linux", "Hosts - Linux"), ("other", "Hosts - Other")):
        ws = wb.create_sheet(title)
        _write_hosts_sheet(ws, split.get(cat, OrderedDict()), hdr_font, hdr_fill, bd, center, left)

    wb.save(out)

def process_one(inp, out=None):
    print(f"\nInput: {inp}")
    data=load_json(inp)
    total=sum(len(h.get("software",[])) for h in data.get("hosts",[]))
    split, host_counts=dedup_by_os(data)
    mw=len(_merge_uniq_by_name(split["windows"])); ml=len(_merge_uniq_by_name(split["linux"])); mo=len(_merge_uniq_by_name(split["other"]))
    print(f"Hosts {len(data.get('hosts',[]))} (W={host_counts['windows']} L={host_counts['linux']} O={host_counts['other']}) total {total} unique software W={mw} L={ml} O={mo}")
    if out is None:
        out = get_out_dir() / f"software_inventory_{Path(inp).stem}.xlsx"
    else:
        out = Path(out)
        if not out.is_absolute():
            out = get_out_dir() / out
        out.parent.mkdir(parents=True, exist_ok=True)
    write_excel(split,str(out),host_counts); print(f"Saved {out}"); return 0

USAGE = """Usage: py json_to_excel.py <input.json> [-o output.xlsx]

Converts one Kaspersky inventory JSON to Excel (split by OS).

Example:
  py json_to_excel.py out\\kaspersky_inventory_20260916_141859.json -o out\\software_inventory.xlsx

Paths can also be driven from gui.py (Convert button)."""

def main(argv=None):
    """Single-file mode only: needs input, optional -o. No args -> usage."""
    import argparse
    ap = argparse.ArgumentParser(description="Convert Kaspersky inventory JSON to Excel (split by OS)")
    ap.add_argument("input", nargs="?", help="input JSON file")
    ap.add_argument("-o", "--output", help="output XLSX path (default: out/software_inventory_<name>.xlsx)")
    a = ap.parse_args(argv)
    if not a.input:
        print(USAGE)
        return 1
    return process_one(a.input, a.output)

if __name__=="__main__":
    raise SystemExit(main())
